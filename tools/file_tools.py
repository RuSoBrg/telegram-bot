import base64
from io import BytesIO
import json
import logging
import os
import re
import time
import zipfile
from html import escape
from typing import Any

from docx import Document
from langchain_core.tools import tool
from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from domain.limits import load_runtime_limits
from domain.object_store import ObjectStore, create_object_store

logger = logging.getLogger(__name__)
store: ObjectStore = create_object_store()


def _filename(filename: str, extension: str, prefix: str) -> str:
    requested = filename.strip() or prefix
    safe = re.sub(r"[^a-zA-Z0-9._-]+", "-", requested).strip(".-_") or prefix
    return f"{safe}-{time.time_ns()}{extension}"


def _validate_content(content: str) -> str:
    if not content or not content.strip():
        raise ValueError("O conteúdo do ficheiro não pode estar vazio.")
    limit = load_runtime_limits().max_generated_content_chars
    if len(content) > limit:
        raise ValueError(f"O conteúdo excede o limite de {limit} caracteres.")
    return content.strip()


def _artifact(filename: str, uri: str, mime_type: str, data: bytes) -> str:
    max_size = int(os.getenv("MAX_ARTIFACT_BYTES", str(50 * 1024 * 1024)))
    if len(data) > max_size:
        raise ValueError(f"O artefacto excede o limite de {max_size} bytes.")
    return json.dumps(
        {
            "tipo": "ficheiro_gerado",
            "filename": filename,
            "key": filename,
            "uri": uri,
            "mime_type": mime_type,
            "size_bytes": len(data),
            "sha256": __import__("hashlib").sha256(data).hexdigest(),
        },
        ensure_ascii=False,
    )


def _publish(filename: str, data: bytes, mime_type: str) -> str:
    uri = store.put_bytes(filename, data)
    return _artifact(filename, uri, mime_type, data)


def _tool_error(error: Exception) -> str:
    logger.exception("Falha ao gerar ficheiro")
    return json.dumps({"tipo": "erro_ficheiro", "erro": str(error)}, ensure_ascii=False)


def _require_env(*names: str) -> dict[str, str]:
    values = {name: os.getenv(name, "").strip() for name in names}
    missing = [name for name, value in values.items() if not value]
    if missing:
        raise RuntimeError(f"Configuração em falta: {', '.join(missing)}.")
    return values


def _google_genai_client():
    from google import genai

    return genai.Client(api_key=_require_env("GOOGLE_API_KEY")["GOOGLE_API_KEY"])


def _inline_bytes(response: Any) -> bytes:
    for candidate in getattr(response, "candidates", []) or []:
        for part in getattr(getattr(candidate, "content", None), "parts", []) or []:
            inline_data = getattr(part, "inline_data", None)
            data = getattr(inline_data, "data", None) if inline_data else None
            if data:
                return data
    raise RuntimeError("O modelo não devolveu dados multimédia.")


def _lines(content: str) -> list[str]:
    return [line.strip() for line in content.splitlines() if line.strip()]


def _read_key(key: str) -> bytes:
    reader = getattr(store, "get_bytes", None)
    if reader is None:
        raise ValueError("Este backend não suporta leitura de artefactos para ZIP.")
    return reader(key)


@tool
def empacotar_artefactos(
    caminhos: list[str], nome_ficheiro: str = "pacote-formacao"
) -> str:
    """Cria um ZIP em memória a partir de chaves de objectos."""
    try:
        if not caminhos:
            raise ValueError("É necessário indicar pelo menos um ficheiro.")
        output = BytesIO()
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
            for key in caminhos:
                bundle.writestr(key.rsplit("/", 1)[-1], _read_key(key))
        data = output.getvalue()
        return _publish(_filename(nome_ficheiro, ".zip", "pacote"), data, "application/zip")
    except (OSError, ValueError, KeyError) as error:
        return _tool_error(error)


@tool
def gerar_documento_word(
    conteudo: str, titulo: str = "", nome_ficheiro: str = "documento"
) -> str:
    try:
        document = Document()
        conteudo = _validate_content(conteudo)
        if titulo.strip():
            document.add_heading(titulo.strip(), level=0)
        for line in _lines(conteudo):
            if line.startswith("### "):
                document.add_heading(line[4:], level=2)
            elif line.startswith("## "):
                document.add_heading(line[3:], level=1)
            elif line.startswith("# "):
                document.add_heading(line[2:], level=0)
            elif line.startswith(("- ", "* ")):
                document.add_paragraph(line[2:], style="List Bullet")
            else:
                document.add_paragraph(line)
        output = BytesIO()
        document.save(output)
        return _publish(
            _filename(nome_ficheiro, ".docx", "documento"),
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    except (OSError, ValueError) as error:
        return _tool_error(error)


@tool
def gerar_apresentacao_powerpoint(
    conteudo: str, titulo: str = "", nome_ficheiro: str = "apresentacao"
) -> str:
    try:
        conteudo = _validate_content(conteudo)
        presentation = Presentation()
        title = titulo.strip() or _lines(conteudo)[0][:120]
        title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        title_slide.shapes.title.text = title
        title_slide.placeholders[1].text = "Gerado pelo assistente"
        for block in (b.strip() for b in re.split(r"\n\s*\n", conteudo) if b.strip()):
            lines = _lines(block)
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = lines[0][:120]
            body = slide.placeholders[1].text_frame
            body.clear()
            for index, line in enumerate(lines[1:] or lines):
                paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
                paragraph.text = line.removeprefix("- ").removeprefix("* ")
        output = BytesIO()
        presentation.save(output)
        return _publish(
            _filename(nome_ficheiro, ".pptx", "apresentacao"),
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
    except (OSError, ValueError) as error:
        return _tool_error(error)


@tool
def gerar_documento_pdf(
    conteudo: str, titulo: str = "", nome_ficheiro: str = "documento"
) -> str:
    try:
        conteudo = _validate_content(conteudo)
        styles = getSampleStyleSheet()
        story: list[Any] = []
        if titulo.strip():
            story.extend([Paragraph(escape(titulo.strip()), styles["Title"]), Spacer(1, 8 * mm)])
        for line in _lines(conteudo):
            style = styles["Heading2"] if line.startswith("#") else styles["BodyText"]
            story.extend(
                [Paragraph(escape(line.lstrip("# ").strip()), style), Spacer(1, 3 * mm)]
            )
        output = BytesIO()
        SimpleDocTemplate(output, pagesize=A4).build(story)
        return _publish(_filename(nome_ficheiro, ".pdf", "documento"), output.getvalue(), "application/pdf")
    except (OSError, ValueError) as error:
        return _tool_error(error)


@tool
def gerar_fotografia(prompt: str, nome_ficheiro: str = "imagem", modelo: str = "") -> str:
    try:
        response = _google_genai_client().models.generate_content(
            model=modelo or os.getenv("GEMINI_IMAGE_MODEL", "gemini-2.5-flash-image"),
            contents=_validate_content(prompt),
            config={"response_modalities": ["TEXT", "IMAGE"]},
        )
        return _publish(_filename(nome_ficheiro, ".png", "imagem"), _inline_bytes(response), "image/png")
    except (OSError, RuntimeError, ValueError) as error:
        return _tool_error(error)


@tool
def gerar_video(
    prompt: str,
    nome_ficheiro: str = "video",
    modelo: str = "",
    duracao_segundos: int = 8,
) -> str:
    try:
        max_seconds = min(8, load_runtime_limits().max_media_seconds)
        if not 1 <= duracao_segundos <= max_seconds:
            raise ValueError(f"A duração do vídeo deve estar entre 1 e {max_seconds} segundos.")
        client = _google_genai_client()
        operation = client.models.generate_videos(
            model=modelo or os.getenv("GEMINI_VIDEO_MODEL", "veo-3.1-generate-preview"),
            prompt=_validate_content(prompt),
            config={"duration_seconds": duracao_segundos, "number_of_videos": 1},
        )
        deadline = time.monotonic() + load_runtime_limits().tool_timeout_seconds
        while not operation.done:
            if time.monotonic() >= deadline:
                raise RuntimeError("A geração do vídeo excedeu o tempo limite configurado.")
            time.sleep(10)
            operation = client.operations.get(operation)
        if getattr(operation, "error", None):
            raise RuntimeError(str(operation.error))
        generated = getattr(getattr(operation, "response", None), "generated_videos", None) or []
        video = getattr(generated[0], "video", None) if generated else None
        data = getattr(video, "video_bytes", None) if video else None
        if not data:
            raise RuntimeError("A resposta do modelo não contém os bytes do vídeo.")
        return _publish(_filename(nome_ficheiro, ".mp4", "video"), data, "video/mp4")
    except (OSError, RuntimeError, ValueError) as error:
        return _tool_error(error)


@tool
def gerar_audio(texto: str, nome_ficheiro: str = "audio", idioma: str = "pt-PT") -> str:
    try:
        _require_env("GOOGLE_CLOUD_PROJECT")
        from google.cloud import texttospeech

        response = texttospeech.TextToSpeechClient().synthesize_speech(
            input=texttospeech.SynthesisInput(text=_validate_content(texto)),
            voice=texttospeech.VoiceSelectionParams(
                language_code=idioma,
                name="pt-PT-Wavenet-A" if idioma == "pt-PT" else "",
            ),
            audio_config=texttospeech.AudioConfig(
                audio_encoding=texttospeech.AudioEncoding.MP3
            ),
        )
        return _publish(_filename(nome_ficheiro, ".mp3", "audio"), response.audio_content, "audio/mpeg")
    except (OSError, RuntimeError, ValueError) as error:
        return _tool_error(error)


@tool
def gerar_musica(
    prompt: str, nome_ficheiro: str = "musica", duracao_segundos: int = 30
) -> str:
    try:
        max_seconds = min(60, load_runtime_limits().max_media_seconds)
        if not 1 <= duracao_segundos <= max_seconds:
            raise ValueError(f"A duração da música deve estar entre 1 e {max_seconds} segundos.")
        config = _require_env("GOOGLE_CLOUD_PROJECT", "GOOGLE_CLOUD_LOCATION")
        import google.auth
        from google.auth.transport.requests import AuthorizedSession

        credentials, _ = google.auth.default()
        response = AuthorizedSession(credentials).post(
            f"https://{config['GOOGLE_CLOUD_LOCATION']}-aiplatform.googleapis.com/v1/"
            f"projects/{config['GOOGLE_CLOUD_PROJECT']}/locations/{config['GOOGLE_CLOUD_LOCATION']}/"
            f"publishers/google/models/{os.getenv('VERTEX_MUSIC_MODEL', 'lyria-002')}:predict",
            json={
                "instances": [{"prompt": _validate_content(prompt)}],
                "parameters": {"duration_seconds": duracao_segundos},
            },
            timeout=load_runtime_limits().tool_timeout_seconds,
        )
        response.raise_for_status()
        predictions = response.json().get("predictions", [])
        encoded = predictions[0].get("audioContent") if predictions else None
        if not encoded:
            raise RuntimeError("O modelo Lyria não devolveu áudio.")
        return _publish(_filename(nome_ficheiro, ".wav", "musica"), base64.b64decode(encoded), "audio/wav")
    except (OSError, RuntimeError, ValueError, IndexError) as error:
        return _tool_error(error)
